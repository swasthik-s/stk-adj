"""
xlsx_template.py — use the store's own Excel template as the output.

The built-in PDF in item_templates.py is a rebuild from screenshots. It is
close, but it is not the real file: the logo, the exact widths and whatever
else lives in the workbook are not in it. This module takes the actual .xlsx
the store already types into, writes the rows into it, and hands it back —
so the output is the template, not an imitation of it.

Three steps, kept separate so each can be checked on its own:

    scan(bytes)                  work out where everything is
    fill(bytes, layout, ...)     write the data in, keeping every style
    to_pdf(bytes)                convert, if LibreOffice is installed

Detection is a starting point, never the last word. Every guess scan() makes
is shown in the app for the person to correct before anything is saved — a
template that is silently mapped wrong produces a sheet that looks right and
says the wrong thing, which is worse than one that plainly fails.
"""

import re
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# ==================== what a column can mean ====================
# Each field key, with the headings seen on the sheets in use. Matching is on
# letters and digits only, so "S.No", "S. No" and "SNO" are the same thing.
COLUMN_ALIASES = {
    "sno": ["sno", "slno", "sl", "serial", "srno", "no"],
    "subcat": ["subcategory", "subcat", "category", "subgroup"],
    "single": ["singlebarcode", "barcode", "itembarcode", "pcsbarcode",
               "singlebc"],
    "outer": ["outerbarcode", "outerbc", "casebarcode", "cartonbarcode"],
    "desc": ["description", "itemname", "itemdescription", "productname",
             "particulars"],
    "old_desc": ["olddescription", "existingdescription", "previousname",
                 "olditemname", "oldname"],
    "new_desc": ["newdescription", "correcteddescription", "newname",
                 "newitemname", "changeto"],
    "unit": ["unit", "uom", "units"],
    "packing": ["packing", "packqty", "pack", "packsize", "qtyperouter",
                "conversion", "conv"],
    "cost": ["cost", "costprice", "purchasecost", "landedcost", "cp"],
    "rsp": ["rsp", "sellingprice", "retailprice", "price", "mrp", "sp"],
    "gp": ["gp", "gppercent", "gp%", "margin", "marginpercent", "gpp"],
}

# Header lines above the table.
FIELD_ALIASES = {
    "purpose": ["purpose"],
    "date": ["date"],
    "vendor": ["vendor", "supplier", "vendorname"],
    "maingrp": ["maingrp", "maingroup", "group", "mainggrp", "grp"],
    "reason": ["reason"],
    "remark": ["remark", "remarks", "note"],
}

FIELD_LABELS_SHORT = {"purpose": "PURPOSE", "date": "DATE",
                      "vendor": "VENDOR", "maingrp": "Main Grp",
                      "reason": "REASON", "remark": "Remark"}

SIGN_ALIASES = {
    "prepared": ["preparedby", "prepared"],
    "purchaser": ["concernedpurchaser", "purchaser", "concerned"],
    "approved": ["approvedby", "approved", "authorisedby", "authorizedby"],
}

NUMERIC = {"cost", "rsp", "gp", "packing"}
# Barcodes must stay text: Excel turns a long number into 1.00071E+10 and
# drops any leading zero.
TEXT_COLS = {"single", "outer", "sno"}

_KEEP = re.compile(r"[^a-z0-9%]")


def norm(s):
    return _KEEP.sub("", str(s or "").strip().lower())


def _match(text, aliases):
    """Exact normalised match first, then 'starts with' — so 'DATE :' finds
    date, but 'DATE OF BIRTH' would not beat a real 'date' cell elsewhere."""
    n = norm(text)
    if not n:
        return None
    for key, names in aliases.items():
        if n in names:
            return key
    for key, names in aliases.items():
        for name in names:
            if n.startswith(name) and len(n) <= len(name) + 3:
                return key
    return None


# ==================== reading the sheet ====================
def _text_grid(ws, max_row, max_col):
    """Cell text by (row, col), merged cells resolved to their top-left value
    so a label inside a merge is still findable."""
    grid = {}
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for c in row:
            if c.value is not None and str(c.value).strip():
                grid[(c.row, c.column)] = str(c.value).strip()
    for rng in ws.merged_cells.ranges:
        top = grid.get((rng.min_row, rng.min_col))
        if top:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    grid.setdefault((r, c), top)
    return grid


def _merged_end(ws, row, col):
    """If a cell is inside a merge, the top-left of that merge — the only cell
    openpyxl will let you write to."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return rng.min_row, rng.min_col
    return row, col


def scan(data, sheet=None):
    """Work out the shape of an uploaded template.

    Returns a layout dict:
        sheet        sheet name used
        header_row   the row holding the column headings
        data_start   first row data goes into
        data_rows    how many ruled blank rows follow before something else
        columns      {field key: column number}
        fields       {field key: (row, col) of the cell the value goes in}
        signs        {prepared|purchaser|approved: (row, col)}
        unmapped     headings that matched nothing, for the person to map
        warnings     anything that needs a human look
    """
    wb = load_workbook(BytesIO(data))
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    max_row = min(ws.max_row or 1, 200)
    max_col = min(ws.max_column or 1, 40)
    grid = _text_grid(ws, max_row, max_col)

    # ---- header row: the row where most cells read like column headings ---
    best, best_hits = None, 0
    for r in range(1, max_row + 1):
        hits = {}
        for c in range(1, max_col + 1):
            key = _match(grid.get((r, c)), COLUMN_ALIASES)
            if key and key not in hits:
                hits[key] = c
        # A row needs a real table's worth of headings, and must include at
        # least one of the two columns every one of these sheets has.
        if len(hits) >= 3 and ({"desc", "new_desc"} & set(hits)
                               or {"cost", "rsp"} <= set(hits)):
            if len(hits) > best_hits:
                best, best_hits = (r, hits), len(hits)

    warnings = []
    if not best:
        return {"sheet": ws.title, "header_row": None, "data_start": None,
                "data_rows": 0, "columns": {}, "fields": {}, "signs": {},
                "unmapped": [], "headings": {},
                "warnings": ["No header row found. Is this the right sheet, "
                             "and does it have a row of column headings such "
                             "as Description, COST, RSP?"]}

    header_row, columns = best
    data_start = header_row + 1

    # ---- every heading on that row, mapped or not -------------------------
    headings, unmapped = {}, []
    for c in range(1, max_col + 1):
        txt = grid.get((header_row, c))
        if not txt:
            continue
        headings[c] = txt
        if _match(txt, COLUMN_ALIASES) is None:
            unmapped.append((c, txt))

    # ---- how many rows the table has room for ----------------------------
    # The ruled box is the real capacity, and it is what the person sees as
    # "the table". Two things make a row look occupied when it is not:
    #
    #   a GP% formula. Every one of these templates ships with
    #   =(I7-H7)/I7*100 already sitting in the blank rows. That is part of
    #   the empty form, not data, and reading it as data made the capacity
    #   come out as zero — so rows were inserted into a table that had room,
    #   the signature block was pushed down, and the untouched formulas below
    #   the last item printed #DIV/0!.
    #
    #   the signature labels. They stop the count, which is the point.
    def ruled(r):
        for c in columns.values():
            b = ws.cell(row=r, column=c).border
            if any(getattr(b, s).style for s in
                   ("left", "right", "top", "bottom")):
                return True
        return False

    def occupied(r):
        """Real content in this row — formulas do not count."""
        for c in range(1, max_col + 1):
            txt = grid.get((r, c))
            if not txt or str(txt).startswith("="):
                continue
            if _match(txt, SIGN_ALIASES):
                return True                 # signature block: stop here
            if c in columns.values():
                return True                 # typed data in a mapped column
        return False

    data_rows = 0
    for r in range(data_start, max_row + 1):
        if occupied(r) or not ruled(r):
            break
        data_rows += 1
        if data_rows > 200:
            break
    if data_rows == 0:                      # template has no ruling at all
        for r in range(data_start, max_row + 1):
            if occupied(r):
                break
            data_rows += 1
            if data_rows > 80:
                break

    # ---- the header lines above the table --------------------------------
    fields = {}
    for r in range(1, header_row):
        for c in range(1, max_col + 1):
            key = _match(grid.get((r, c)), FIELD_ALIASES)
            if not key or key in fields:
                continue
            # The value sits in the first empty cell to the right of the
            # label. Labels are often "VENDOR :" in one cell, value in the
            # next; sometimes the colon is its own cell.
            for step in range(1, 5):
                tr, tc = _merged_end(ws, r, c + step)
                if (tr, tc) == _merged_end(ws, r, c):
                    continue
                txt = grid.get((r, c + step), "")
                if txt.strip() in (":", "-"):
                    continue
                fields[key] = (tr, tc)
                break

    # ---- signature block --------------------------------------------------
    signs = {}
    for (r, c), txt in list(grid.items()):
        key = _match(txt, SIGN_ALIASES)
        if key and key not in signs:
            signs[key] = _merged_end(ws, r, c)

    # A field whose row is hidden will be written and never seen. The store's
    # creation sheet hides its Remark row, so this is not hypothetical.
    hidden = sorted({FIELD_LABELS_SHORT.get(k, k) for k, (r, _c) in
                     fields.items()
                     if ws.row_dimensions[r].hidden})
    if hidden:
        warnings.append(
            f"{', '.join(hidden)} sits on a hidden row in this template, so "
            f"anything entered there will not print. Unhide the row in Excel "
            f"if you want it shown.")

    if "date" not in fields:
        warnings.append("No DATE cell found — set it below or the date will "
                        "not be printed.")
    if not signs:
        warnings.append("No signature lines found (Prepared By / Concerned "
                        "Purchaser / Approved By). Names will be left off.")
    if data_rows == 0:
        warnings.append("No blank rows under the headings — rows will be "
                        "inserted, which can shift anything below the table.")
    if unmapped:
        warnings.append(f"{len(unmapped)} heading(s) were not recognised: "
                        + ", ".join(t for _c, t in unmapped[:6])
                        + ". Map them below, or they will be left empty.")

    return {"sheet": ws.title, "header_row": header_row,
            "data_start": data_start, "data_rows": data_rows,
            "columns": columns, "fields": fields, "signs": signs,
            "unmapped": unmapped, "headings": headings,
            "warnings": warnings}


# ==================== writing into it ====================
def _copy_style(src, dst):
    if src.has_style:
        dst._style = src._style


def _coerce(key, value):
    """A figure for the number columns, text for the barcodes, so Excel does
    not reformat a 13-digit barcode into scientific notation."""
    if value is None:
        return None
    if key in TEXT_COLS:
        s = str(value).strip()
        return s or None
    if key in NUMERIC:
        try:
            f = float(str(value).replace(",", "").strip())
            return None if f != f else f
        except (TypeError, ValueError):
            return None
    s = str(value).strip()
    return s or None


def fill(data, layout, header, rows, *, gp_values=None, blanks=None,
         page_fit=True):
    """Return the template with the data written in.

    layout   as returned by scan(), possibly corrected by the person
    header   date, vendor, maingrp, reason, remark, prepared, purchaser…
    rows     list of dicts keyed by field key
    gp_values  GP% per row, worked out by the caller. Ignored when the
             template's own GP% cell holds a formula — that formula is the
             store's, and it is copied down rather than overwritten.
    blanks   what an empty cell prints as, per field key, e.g. the
             NEED BARCODE the creation sheets use instead of a gap.
    page_fit force the sheet onto one page across when the template does not
             say how it should print. Off leaves the workbook untouched.
    """
    blanks = blanks or {}
    wb = load_workbook(BytesIO(data))
    ws = wb[layout["sheet"]] if layout.get("sheet") in wb.sheetnames \
        else wb[wb.sheetnames[0]]

    columns = {k: int(v) for k, v in (layout.get("columns") or {}).items()}
    start = int(layout["data_start"])
    spare = int(layout.get("data_rows") or 0)
    n = len(rows)

    # ---- make room FIRST --------------------------------------------------
    # Order matters. These templates put a merged signature block under the
    # table (B15:C16, H15:J16 and so on). openpyxl's insert_rows moves cell
    # values but leaves merged ranges and row heights where they were, so the
    # merges end up lying across the new data rows — writes into them are
    # silently dropped and the signature block vanishes. So the sheet is
    # stretched before anything is written, and the merges and heights are
    # carried down by hand.
    extra, at = 0, None
    if n > spare:
        extra = n - spare
        at = start + max(spare, 1)

        moved = [(r.min_row, r.min_col, r.max_row, r.max_col)
                 for r in list(ws.merged_cells.ranges) if r.min_row >= at]
        for r1, c1, r2, c2 in moved:
            ws.unmerge_cells(start_row=r1, start_column=c1,
                             end_row=r2, end_column=c2)

        heights = {r: d.height for r, d in ws.row_dimensions.items()
                   if r >= at and d.height is not None}
        hidden = {r for r, d in ws.row_dimensions.items()
                  if r >= at and d.hidden}

        ws.insert_rows(at, extra)

        for r1, c1, r2, c2 in moved:
            ws.merge_cells(start_row=r1 + extra, start_column=c1,
                           end_row=r2 + extra, end_column=c2)
        for r, h in sorted(heights.items(), reverse=True):
            ws.row_dimensions[r + extra].height = h
        for r in sorted(hidden, reverse=True):
            ws.row_dimensions[r + extra].hidden = True
            ws.row_dimensions[r].hidden = False

        # New rows arrive unstyled — copy the ruling off the first data row
        # so the table keeps its box.
        for i in range(extra):
            for c in range(1, (ws.max_column or 1) + 1):
                _copy_style(ws.cell(row=start, column=c),
                            ws.cell(row=at + i, column=c))

    def moved_row(r):
        """Where a row ended up once the sheet was stretched."""
        r = int(r)
        return r + extra if (at is not None and r >= at) else r

    # ---- header lines -----------------------------------------------------
    for key, pos in (layout.get("fields") or {}).items():
        if key == "purpose":
            continue                      # the template already says it
        val = header.get(key)
        if val in (None, ""):
            continue
        ws.cell(row=moved_row(pos[0]), column=int(pos[1])).value = val

    for key, pos in (layout.get("signs") or {}).items():
        name = header.get(key)
        if not name:
            continue
        r, c = moved_row(pos[0]), int(pos[1])
        cur = ws.cell(row=r, column=c).value
        label = str(cur).strip() if cur else ""
        # Keep the printed label and append the name, rather than replacing
        # "Prepared By :" with a bare name.
        ws.cell(row=r, column=c).value = (
            f"{label} {name}".strip() if label.rstrip().endswith(":")
            else (f"{label}: {name}" if label else name))

    # ---- is GP% a formula in the template? --------------------------------
    gp_col = columns.get("gp")
    gp_formula = None
    if gp_col:
        v = ws.cell(row=start, column=gp_col).value
        if isinstance(v, str) and v.startswith("="):
            gp_formula = v

    # ---- the rows ---------------------------------------------------------
    for i, row in enumerate(rows):
        r = start + i
        for key, col in columns.items():
            if key == "sno":
                ws.cell(row=r, column=col).value = i + 1
                continue
            if key == "gp":
                if gp_formula is not None:
                    ws.cell(row=r, column=col).value = _shift(gp_formula,
                                                              start, r)
                elif gp_values is not None and i < len(gp_values):
                    g = gp_values[i]
                    ws.cell(row=r, column=col).value = (
                        None if g is None else round(float(g), 2))
                continue
            v = _coerce(key, row.get(key))
            if v is None and key in blanks:
                v = blanks[key]
            ws.cell(row=r, column=col).value = v

    # ---- clear the rest of the band ---------------------------------------
    # Every template row past the last item, across the whole width, not just
    # the mapped columns. The GP% formulas that ship in the blank rows live
    # here; left alone they print #DIV/0! under the last real line.
    for r in range(start + n, start + max(spare, n)):
        for col in range(1, (ws.max_column or 1) + 1):
            ws.cell(row=r, column=col).value = None

    if page_fit:
        _fit_page(ws, last_row=start + max(spare, n))

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _fit_page(ws, last_row=None):
    """Make the sheet print on one page across.

    A template that has never been printed carries no page setup, and
    LibreOffice then breaks the columns over two or three pages — the sheet
    is still correct but nobody can sign it. Anything the template already
    specifies is left alone; only the gaps are filled.
    """
    ps = ws.page_setup
    props = ws.sheet_properties.pageSetUpPr
    if not ps.orientation:
        ps.orientation = "landscape"
    if not ps.paperSize:
        ps.paperSize = ws.PAPERSIZE_A4
    already = bool(props and props.fitToPage)
    if not already:
        from openpyxl.worksheet.properties import PageSetupProperties
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ps.fitToWidth = 1
        ps.fitToHeight = 0          # as many pages down as it needs
    if not ws.print_area:
        # To the last used row, never just the end of the table: the
        # signature block sits below it, and a print area that stops at the
        # table silently produces an approval sheet nobody can sign.
        end_row = max(int(ws.max_row or 1), int(last_row or 1))
        end_col = get_column_letter(max(ws.max_column or 1, 1))
        ws.print_area = f"A1:{end_col}{end_row}"


_REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def _shift(formula, from_row, to_row):
    """Move a relative formula down the sheet. =(I10-H10)/I10 at row 10
    becomes =(I11-H11)/I11 at row 11. Absolute rows ($10) are left alone."""
    delta = to_row - from_row
    if not delta:
        return formula

    def sub(m):
        dollar_col, col, dollar_row, row = m.groups()
        if dollar_row:
            return m.group(0)
        return f"{dollar_col}{col}{dollar_row}{int(row) + delta}"

    return _REF.sub(sub, formula)


# ==================== xlsx -> pdf ====================
def converter():
    """Path to LibreOffice, or None. The app shows this, so the person knows
    up front whether a PDF is possible rather than after pressing Generate."""
    return shutil.which("soffice") or shutil.which("libreoffice")


def to_pdf(xlsx_bytes, timeout=120):
    """(pdf bytes, None) or (None, why). Never raises — a failed conversion
    must still leave the person with their .xlsx."""
    exe = converter()
    if not exe:
        return None, ("LibreOffice is not installed on this server, so the "
                      "sheet can only be given as .xlsx. Add `libreoffice-calc` "
                      "to packages.txt and reboot the app to turn PDFs on.")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "sheet.xlsx"
        src.write_bytes(xlsx_bytes)
        try:
            p = subprocess.run(
                [exe, "--headless", "--norestore", "--convert-to", "pdf",
                 "--outdir", tmp, str(src)],
                capture_output=True, timeout=timeout,
                env={"HOME": tmp, "PATH": "/usr/bin:/bin"})
        except subprocess.TimeoutExpired:
            return None, f"LibreOffice did not finish within {timeout}s."
        except Exception as e:
            return None, f"{type(e).__name__}: {e}"
        pdf = Path(tmp) / "sheet.pdf"
        if not pdf.exists():
            err = (p.stderr or b"").decode("utf8", "replace").strip()
            return None, f"LibreOffice produced no PDF. {err[:300]}"
        return pdf.read_bytes(), None


def sheet_names(data):
    try:
        return load_workbook(BytesIO(data), read_only=True).sheetnames
    except Exception:
        return []
